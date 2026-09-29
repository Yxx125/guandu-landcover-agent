"""脚本名称：index_documents.py

将 documents/ 中的 UTF-8 文本、PDF 和 Excel (.xlsx) 切块、向量化并持久化到本地 Chroma。
文件缺失会直接报错；不制造虚构的规划文本或研究资料。
"""

import hashlib
import sys
from pathlib import Path

import chromadb
from openpyxl import load_workbook
from pypdf import PdfReader
from sentence_transformers import SentenceTransformer

PROJECT_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_DIR))
from config import EMBEDDING_MODEL, CHROMA_DIR, DOCUMENTS_DIR  # noqa: E402

DOCUMENT_DIR = DOCUMENTS_DIR
CLCD_ZH_NAMES = {
    1: "耕地", 2: "林地", 3: "灌丛", 4: "草地", 5: "水体",
    6: "积雪／冰川", 7: "裸地", 8: "建设用地（不透水面）", 9: "湿地",
}


def read_document(path: Path) -> str:
    if path.suffix.lower() == ".txt":
        return path.read_text(encoding="utf-8-sig")
    if path.suffix.lower() == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if path.suffix.lower() == ".xlsx":
        lines = []
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook:
                rows = sheet.iter_rows(values_only=True)
                header = next(rows, None)
                if header is None:
                    continue
                columns = [str(value).strip() if value is not None else f"第{i + 1}列"
                           for i, value in enumerate(header)]
                for row_number, row in enumerate(rows, start=2):
                    cells = [f"{name}: {value}" for name, value in zip(columns, row)
                             if value is not None and str(value).strip()]
                    if cells:
                        if {"ID", "Class", "Color"}.issubset(columns):
                            try:
                                code = int(row[columns.index("ID")])
                            except (TypeError, ValueError):
                                code = None
                            if code in CLCD_ZH_NAMES:
                                cells.append(f"系统对应中文类别: {CLCD_ZH_NAMES[code]}")
                        lines.append(f"文件 {path.name}；工作表 {sheet.title}；"
                                     f"第 {row_number} 行；" + "；".join(cells))
        finally:
            workbook.close()
        return "\n".join(lines)
    raise ValueError(f"不支持的文档类型：{path.name}")


def chunks(text: str, length: int = 500, overlap: int = 80):
    cleaned = " ".join(text.split())
    step = length - overlap
    for position in range(0, len(cleaned), step):
        value = cleaned[position:position + length]
        if len(value) >= 40:
            yield value
        if position + length >= len(cleaned):
            break


def main():
    paths = sorted(
        path for path in DOCUMENT_DIR.rglob("*")
        if path.is_file() and not path.name.startswith("~$")
        and path.suffix.lower() in {".txt", ".pdf", ".xlsx"}
    ) if DOCUMENT_DIR.exists() else []
    if not paths:
        raise SystemExit("documents/ 中没有 .txt、.pdf 或 .xlsx 文件")
    ids, texts, metadata = [], [], []
    for path in paths:
        for index, text in enumerate(chunks(read_document(path))):
            key = f"{path.relative_to(DOCUMENT_DIR)}:{index}"
            ids.append(hashlib.sha256(key.encode("utf-8")).hexdigest())
            texts.append(text)
            metadata.append({"source": str(path.relative_to(DOCUMENT_DIR)), "chunk": index})
    if not texts:
        raise SystemExit("文档没有可用文本；扫描版 PDF 请先进行 OCR 再导入")
    print(f"读取 {len(paths)} 份文档、{len(texts)} 个文本块；正在加载本地向量模型……")
    model = SentenceTransformer(EMBEDDING_MODEL)
    client = chromadb.PersistentClient(path=str(CHROMA_DIR))
    collection = client.get_or_create_collection(
        "guandu_documents", configuration={"hnsw": {"space": "cosine"}}
    )
    for offset in range(0, len(texts), 32):
        batch = texts[offset:offset + 32]
        vectors = model.encode(
            batch, normalize_embeddings=True, show_progress_bar=False
        ).tolist()
        collection.upsert(
            ids=ids[offset:offset + 32],
            documents=batch,
            metadatas=metadata[offset:offset + 32],
            embeddings=vectors,
        )
        print(f"已入库 {min(offset + len(batch), len(texts))}/{len(texts)}")
    # 重建时删除上次入库而本次已移走的文件/旧切块，避免检索到过期资料。
    current_ids = set(ids)
    previous_ids = collection.get(include=[])['ids']
    stale_ids = [item for item in previous_ids if item not in current_ids]
    for offset in range(0, len(stale_ids), 500):
        collection.delete(ids=stale_ids[offset:offset + 500])
    print(f"Chroma 导入完成：{collection.count()} 个文本块。")


if __name__ == "__main__":
    main()
